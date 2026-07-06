from flask import Flask, render_template, request, jsonify, send_file, redirect, url_for, flash, make_response
from flask_login import LoginManager, UserMixin, login_user, logout_user, login_required, current_user
from flask_bcrypt import Bcrypt
import os
import re
import io
import json
import pdfplumber
import anthropic
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable, Table, TableStyle
from reportlab.lib import colors
from datetime import datetime
import time

try:
    from docx import Document as DocxDocument
    DOCX_SUPPORT = True
except ImportError:
    DOCX_SUPPORT = False

try:
    import requests as http_requests
    from bs4 import BeautifulSoup
    SCRAPE_SUPPORT = True
except ImportError:
    SCRAPE_SUPPORT = False

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'dev-only-change-in-production')

login_manager = LoginManager(app)
login_manager.login_view = 'login'
login_manager.login_message = 'Bitte melde dich an, um fortzufahren.'
login_manager.login_message_category = 'info'
bcrypt = Bcrypt(app)

BASE_DIR     = os.path.dirname(os.path.abspath(__file__))
UPLOAD_FOLDER = os.path.join(BASE_DIR, 'uploads')
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

# ── DB Driver (PostgreSQL on Railway, SQLite for local dev) ───────────────────
DATABASE_URL = os.environ.get('DATABASE_URL')   # Railway injects this automatically
_PG          = bool(DATABASE_URL)

if _PG:
    import psycopg2
    import psycopg2.extras
else:
    import sqlite3
    DATABASE = os.path.join(BASE_DIR, 'bewerbungen.db')


class _Conn:
    """Thin wrapper so psycopg2 and sqlite3 share one call interface.

    – sqlite3  uses ? placeholders and conn.row_factory for dict-like rows
    – psycopg2 uses %s placeholders and RealDictCursor for dict-like rows
    Both return objects where row['column'] works, so all templates stay unchanged.
    """
    def __init__(self, raw):
        self._raw = raw

    def execute(self, sql, params=()):
        if _PG:
            cur = self._raw.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
            cur.execute(sql.replace('?', '%s'), params)
        else:
            cur = self._raw.cursor()
            cur.execute(sql, params)
        return cur

    def commit(self): self._raw.commit()
    def close(self):  self._raw.close()


# ── Auth ──────────────────────────────────────────────────────────────────────

class User(UserMixin):
    def __init__(self, id, email, name, plan='free'):
        self.id = id
        self.email = email
        self.name = name
        self.plan = plan


@login_manager.user_loader
def load_user(user_id):
    conn = get_db()
    row = conn.execute('SELECT id, email, name, plan FROM users WHERE id = ?', (int(user_id),)).fetchone()
    conn.close()
    if row:
        return User(row['id'], row['email'], row['name'], row['plan'])
    return None


STATUS_OPTIONS = [
    'Entwurf', 'Gesendet', 'Telefoninterview',
    'Vorstellungsgespräch', 'Angebot erhalten', 'Angenommen', 'Abgelehnt'
]
STATUS_CSS = {
    'Entwurf':             'draft',
    'Gesendet':            'applied',
    'Telefoninterview':    'phone',
    'Vorstellungsgespräch':'interview',
    'Angebot erhalten':    'offer',
    'Angenommen':          'accepted',
    'Abgelehnt':           'rejected',
}


# ── Database ──────────────────────────────────────────────────────────────────

def get_db():
    if _PG:
        return _Conn(psycopg2.connect(DATABASE_URL))
    raw = sqlite3.connect(DATABASE)
    raw.row_factory = sqlite3.Row
    return _Conn(raw)


def init_db():
    if _PG:
        for attempt in range(10):
            try:
                psycopg2.connect(DATABASE_URL).close()
                break
            except psycopg2.OperationalError:
                if attempt == 9:
                    raise
                time.sleep(3)
    conn = get_db()
    if _PG:
        conn.execute('''CREATE TABLE IF NOT EXISTS users (
            id SERIAL PRIMARY KEY,
            email TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL,
            name TEXT DEFAULT '',
            plan TEXT DEFAULT 'free',
            created_at TEXT DEFAULT CURRENT_DATE::TEXT
        )''')
        conn.execute('''CREATE TABLE IF NOT EXISTS bewerbungen (
            id SERIAL PRIMARY KEY,
            user_id INTEGER REFERENCES users(id) ON DELETE CASCADE,
            firma TEXT NOT NULL DEFAULT '',
            stelle TEXT DEFAULT '',
            url TEXT DEFAULT '',
            stellenanzeige TEXT DEFAULT '',
            anschreiben TEXT DEFAULT '',
            status TEXT DEFAULT 'Entwurf',
            datum_erstellt TEXT DEFAULT CURRENT_DATE::TEXT,
            datum_gesendet TEXT,
            notizen TEXT DEFAULT ''
        )''')
        conn.execute('ALTER TABLE bewerbungen ADD COLUMN IF NOT EXISTS user_id INTEGER REFERENCES users(id) ON DELETE CASCADE')
        # Migrate settings: old schema had key as sole PK, no user_id
        needs_settings_migration = not conn.execute(
            "SELECT 1 FROM information_schema.columns WHERE table_name=? AND column_name=?",
            ('settings', 'user_id')
        ).fetchone()
        if needs_settings_migration:
            conn.execute('DROP TABLE IF EXISTS settings CASCADE')
        conn.execute('''CREATE TABLE IF NOT EXISTS settings (
            user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
            key TEXT NOT NULL,
            value TEXT DEFAULT '',
            PRIMARY KEY (user_id, key)
        )''')
        conn.execute('''CREATE TABLE IF NOT EXISTS chat_nachrichten (
            id SERIAL PRIMARY KEY,
            bewerbung_id INTEGER NOT NULL,
            rolle TEXT NOT NULL,
            inhalt TEXT NOT NULL,
            erstellt_am TEXT DEFAULT TO_CHAR(CURRENT_TIMESTAMP, 'YYYY-MM-DD HH24:MI:SS'),
            FOREIGN KEY (bewerbung_id) REFERENCES bewerbungen(id) ON DELETE CASCADE
        )''')
        # Migrate cv_data: old schema was singleton (id=1), new is per-user
        needs_cv_migration = not conn.execute(
            "SELECT 1 FROM information_schema.columns WHERE table_name=? AND column_name=?",
            ('cv_data', 'user_id')
        ).fetchone()
        if needs_cv_migration:
            conn.execute('DROP TABLE IF EXISTS cv_data CASCADE')
        conn.execute('''CREATE TABLE IF NOT EXISTS cv_data (
            user_id INTEGER PRIMARY KEY REFERENCES users(id) ON DELETE CASCADE,
            data TEXT DEFAULT '{}',
            template TEXT DEFAULT 'klassisch',
            updated_at TEXT DEFAULT TO_CHAR(CURRENT_TIMESTAMP, 'YYYY-MM-DD HH24:MI:SS')
        )''')
    else:
        conn.execute('''CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            email TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL,
            name TEXT DEFAULT '',
            plan TEXT DEFAULT 'free',
            created_at TEXT DEFAULT (strftime('%Y-%m-%d', 'now'))
        )''')
        conn.execute('''CREATE TABLE IF NOT EXISTS bewerbungen (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER REFERENCES users(id) ON DELETE CASCADE,
            firma TEXT NOT NULL DEFAULT '',
            stelle TEXT DEFAULT '',
            url TEXT DEFAULT '',
            stellenanzeige TEXT DEFAULT '',
            anschreiben TEXT DEFAULT '',
            status TEXT DEFAULT 'Entwurf',
            datum_erstellt TEXT DEFAULT (strftime('%Y-%m-%d', 'now')),
            datum_gesendet TEXT,
            notizen TEXT DEFAULT ''
        )''')
        cur = conn.execute("SELECT 1 FROM pragma_table_info('bewerbungen') WHERE name='user_id'")
        if not cur.fetchone():
            conn.execute('ALTER TABLE bewerbungen ADD COLUMN user_id INTEGER REFERENCES users(id) ON DELETE CASCADE')
        # Migrate settings
        cur = conn.execute("SELECT 1 FROM pragma_table_info('settings') WHERE name='user_id'")
        if not cur.fetchone():
            conn.execute('DROP TABLE IF EXISTS settings')
        conn.execute('''CREATE TABLE IF NOT EXISTS settings (
            user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
            key TEXT NOT NULL,
            value TEXT DEFAULT '',
            PRIMARY KEY (user_id, key)
        )''')
        conn.execute('''CREATE TABLE IF NOT EXISTS chat_nachrichten (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            bewerbung_id INTEGER NOT NULL,
            rolle TEXT NOT NULL,
            inhalt TEXT NOT NULL,
            erstellt_am TEXT DEFAULT (datetime('now')),
            FOREIGN KEY (bewerbung_id) REFERENCES bewerbungen(id) ON DELETE CASCADE
        )''')
        # Migrate cv_data
        cur = conn.execute("SELECT 1 FROM pragma_table_info('cv_data') WHERE name='user_id'")
        if not cur.fetchone():
            conn.execute('DROP TABLE IF EXISTS cv_data')
        conn.execute('''CREATE TABLE IF NOT EXISTS cv_data (
            user_id INTEGER PRIMARY KEY REFERENCES users(id) ON DELETE CASCADE,
            data TEXT DEFAULT '{}',
            template TEXT DEFAULT 'klassisch',
            updated_at TEXT DEFAULT (datetime('now'))
        )''')
    # Assign orphaned rows (user_id IS NULL) to the first user.
    # Covers bewerbungen created before auth was added.
    conn.execute('''
        UPDATE bewerbungen SET user_id = (SELECT id FROM users ORDER BY id LIMIT 1)
        WHERE user_id IS NULL AND EXISTS (SELECT 1 FROM users)
    ''')
    conn.commit()
    conn.close()


init_db()


def get_setting(key, default=''):
    if not current_user.is_authenticated:
        return default
    uid = current_user.id
    conn = get_db()
    row = conn.execute('SELECT value FROM settings WHERE user_id = ? AND key = ?', (uid, key)).fetchone()
    conn.close()
    return row['value'] if row and row['value'] else default


def set_setting(key, value):
    if not current_user.is_authenticated:
        return
    uid = current_user.id
    conn = get_db()
    if _PG:
        conn.execute(
            'INSERT INTO settings (user_id, key, value) VALUES (?, ?, ?)'
            ' ON CONFLICT (user_id, key) DO UPDATE SET value = EXCLUDED.value',
            (uid, key, value)
        )
    else:
        conn.execute('INSERT OR REPLACE INTO settings (user_id, key, value) VALUES (?, ?, ?)', (uid, key, value))
    conn.commit()
    conn.close()


# ── Context Processor (global template vars) ─────────────────────────────────

@app.context_processor
def inject_globals():
    if not current_user.is_authenticated:
        return {'nav_total': 0, 'profile_name': '', 'profile_email': '',
                'status_css': STATUS_CSS, 'status_options': STATUS_OPTIONS}
    conn = get_db()
    total = conn.execute('SELECT COUNT(*) as n FROM bewerbungen WHERE user_id = ?', (current_user.id,)).fetchone()['n']
    conn.close()
    return {
        'nav_total': total,
        'profile_name': current_user.name,
        'profile_email': current_user.email,
        'status_css': STATUS_CSS,
        'status_options': STATUS_OPTIONS,
    }


# ── Seiten ────────────────────────────────────────────────────────────────────

@app.route('/landing')
def landing():
    return render_template('landing.html')


@app.route('/jobs')
@login_required
def jobs():
    has_adzuna = bool(
        os.environ.get('ADZUNA_APP_ID') or get_setting('adzuna_app_id')
    )
    return render_template('jobs.html', has_adzuna=has_adzuna)


@app.route('/')
@login_required
def index():
    conn = get_db()
    bewerbungen = conn.execute(
        'SELECT * FROM bewerbungen WHERE user_id = ? ORDER BY datum_erstellt DESC',
        (current_user.id,)
    ).fetchall()
    conn.close()
    stats = {
        'total': len(bewerbungen),
        'entwurf': sum(1 for b in bewerbungen if b['status'] == 'Entwurf'),
        'gesendet': sum(1 for b in bewerbungen if b['status'] in (
            'Gesendet', 'Telefoninterview', 'Vorstellungsgespräch', 'Angebot erhalten')),
        'angenommen': sum(1 for b in bewerbungen if b['status'] == 'Angenommen'),
        'abgelehnt':  sum(1 for b in bewerbungen if b['status'] == 'Abgelehnt'),
    }
    resp = make_response(render_template('index.html', bewerbungen=bewerbungen, stats=stats))
    resp.headers['Cache-Control'] = 'no-store'
    return resp


@app.route('/neue-bewerbung')
@login_required
def neue_bewerbung():
    has_cv  = bool(get_setting('cv_text'))
    has_api = bool(os.environ.get('ANTHROPIC_API_KEY'))
    return render_template('neue_bewerbung.html', has_cv=has_cv, has_api=has_api)


@app.route('/bewerbung/<int:bid>')
@login_required
def bewerbung(bid):
    conn = get_db()
    b = conn.execute('SELECT * FROM bewerbungen WHERE id = ? AND user_id = ?', (bid, current_user.id)).fetchone()
    if not b:
        return redirect(url_for('index'))
    chat = conn.execute(
        'SELECT * FROM chat_nachrichten WHERE bewerbung_id = ? ORDER BY erstellt_am',
        (bid,)
    ).fetchall()
    conn.close()
    return render_template('bewerbung.html', b=b, chat=chat)


@app.route('/lebenslauf')
@login_required
def lebenslauf():
    cv_json, cv_template = _get_cv_json()
    return render_template('lebenslauf.html',
                           cv_text=get_setting('cv_text'),
                           cv_filename=get_setting('cv_filename'),
                           cv_json=cv_json,
                           cv_template=cv_template)


@app.route('/profil', methods=['GET', 'POST'])
@login_required
def profil():
    if request.method == 'POST':
        for field in ['name', 'email', 'telefon', 'strasse', 'plz_ort', 'ort']:
            if field in request.form:
                set_setting(field, request.form[field].strip())
        flash('Profil gespeichert!', 'success')
        return redirect(url_for('profil'))
    s = {k: get_setting(k) for k in ['name', 'email', 'telefon', 'strasse', 'plz_ort', 'ort']}
    return render_template('profil.html', s=s)


@app.route('/einstellungen', methods=['GET', 'POST'])
@login_required
def einstellungen():
    if request.method == 'POST':
        for field in ['name', 'email', 'telefon', 'strasse', 'plz_ort', 'ort']:
            if field in request.form:
                set_setting(field, request.form[field].strip())
        flash('Einstellungen gespeichert!', 'success')
        return redirect(url_for('einstellungen'))
    s = {k: get_setting(k) for k in
         ['name', 'email', 'telefon', 'strasse', 'plz_ort', 'ort', 'cv_filename', 'cv_text']}
    return render_template('einstellungen.html', s=s)


# ── API: Lebenslauf ───────────────────────────────────────────────────────────

@app.route('/api/upload-cv', methods=['POST'])
@login_required
def upload_cv():
    if 'cv' not in request.files:
        return jsonify({'error': 'Keine Datei ausgewählt'}), 400

    file = request.files['cv']
    fname = file.filename.lower()

    if not (fname.endswith('.pdf') or fname.endswith('.docx')):
        return jsonify({'error': 'Nur PDF oder DOCX erlaubt'}), 400

    ext = '.pdf' if fname.endswith('.pdf') else '.docx'
    filepath = os.path.join(UPLOAD_FOLDER, f'cv{ext}')
    file.save(filepath)

    try:
        text = ''
        if ext == '.pdf':
            with pdfplumber.open(filepath) as pdf:
                for page in pdf.pages:
                    text += (page.extract_text() or '') + '\n'
        else:
            if not DOCX_SUPPORT:
                return jsonify({'error': 'DOCX-Unterstützung nicht verfügbar'}), 500
            doc = DocxDocument(filepath)
            text = '\n'.join(p.text for p in doc.paragraphs if p.text.strip())

        text = text.strip()
        set_setting('cv_text', text)
        set_setting('cv_filename', file.filename)
        return jsonify({'success': True, 'filename': file.filename, 'chars': len(text)})
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/api/delete-cv', methods=['POST'])
@login_required
def delete_cv():
    set_setting('cv_text', '')
    set_setting('cv_filename', '')
    for ext in ('.pdf', '.docx'):
        p = os.path.join(UPLOAD_FOLDER, f'cv{ext}')
        if os.path.exists(p):
            os.remove(p)
    return jsonify({'success': True})


# ── API: URL-Extraktion ───────────────────────────────────────────────────────

@app.route('/api/extract-url', methods=['POST'])
@login_required
def extract_url():
    data = request.get_json()
    url = (data.get('url') or '').strip()

    if not url:
        return jsonify({'error': 'URL fehlt'}), 400
    if not SCRAPE_SUPPORT:
        return jsonify({'error': 'URL-Extraktion nicht verfügbar (requests/bs4 fehlen)'}), 500

    try:
        headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'}
        resp = http_requests.get(url, headers=headers, timeout=12)
        resp.raise_for_status()

        soup = BeautifulSoup(resp.text, 'html.parser')
        for tag in soup(['nav', 'header', 'footer', 'script', 'style', 'aside', 'iframe']):
            tag.decompose()

        title = (soup.find('title') or soup.find('h1'))
        title_text = title.get_text(strip=True) if title else ''

        raw = soup.get_text(separator='\n', strip=True)
        raw = re.sub(r'\n{3,}', '\n\n', raw)

        return jsonify({'success': True, 'text': raw[:6000], 'title': title_text})
    except Exception as e:
        return jsonify({'error': f'Seite konnte nicht geladen werden: {str(e)}'}), 500


# ── API: Job-Suche ────────────────────────────────────────────────────────────

COMMON_JOB_TITLES = [
    # IT / Tech
    'Softwareentwickler', 'Software Engineer', 'Frontend Developer', 'Backend Developer',
    'Full Stack Developer', 'Web Developer', 'DevOps Engineer', 'Cloud Engineer',
    'Data Engineer', 'Data Analyst', 'Data Scientist', 'Machine Learning Engineer',
    'Systemadministrator', 'IT-Administrator', 'IT-Support', 'Netzwerkadministrator',
    'IT-Projektmanager', 'Scrum Master', 'Product Owner', 'UX Designer', 'UI Designer',
    'IT-Sicherheitsexperte', 'Cyber Security', 'Fachinformatiker', 'Informatiker',
    # Business / Management
    'Projektmanager', 'Projektleiter', 'Business Analyst', 'Unternehmensberater',
    'Consultant', 'Manager', 'Teamleiter', 'Abteilungsleiter', 'Geschäftsführer',
    'Strategieberater', 'Change Manager', 'Prozessmanager',
    # Marketing / Sales
    'Marketing Manager', 'Online Marketing Manager', 'SEO Manager', 'Social Media Manager',
    'Content Manager', 'Vertriebsmitarbeiter', 'Sales Manager', 'Account Manager',
    'Key Account Manager', 'Außendienstmitarbeiter', 'Marketingleiter',
    # Finance / Accounting
    'Buchhalter', 'Finanzbuchhalter', 'Controller', 'Steuerberater', 'Wirtschaftsprüfer',
    'Finanzanalyst', 'CFO', 'Kreditanalyst', 'Revisor',
    # HR
    'HR Manager', 'Personalreferent', 'Recruiter', 'HR Business Partner',
    'Personalleiter', 'Talent Acquisition',
    # Engineering
    'Ingenieur', 'Maschinenbauingenieur', 'Elektroingenieur', 'Bauingenieur',
    'Mechatroniker', 'Konstrukteur', 'Produktionsingenieur', 'Qualitätsingenieur',
    'Verfahrenstechniker', 'Chemieingenieur', 'Umweltingenieur',
    # Healthcare
    'Krankenpfleger', 'Pflegefachkraft', 'Arzt', 'Zahnarzt', 'Physiotherapeut',
    'Ergotherapeut', 'Rettungssanitäter', 'Medizinischer Fachangestellter',
    # Trades
    'Elektriker', 'Elektroniker', 'Schlosser', 'Mechatroniker', 'Klempner',
    'Tischler', 'Schweißer', 'KFZ-Mechatroniker', 'Lagerlogistiker', 'Fahrer',
    # Education
    'Lehrer', 'Erzieher', 'Sozialpädagoge', 'Trainer', 'Coach',
    # Admin / Office
    'Sachbearbeiter', 'Assistenz', 'Sekretärin', 'Büromanager', 'Empfangsmitarbeiter',
    'Kaufmann', 'Kauffrau', 'Verwaltungsfachangestellter',
]


@app.route('/api/jobs/search')
@login_required
def jobs_search():
    if not SCRAPE_SUPPORT:
        return jsonify({'error': 'requests-Bibliothek fehlt'}), 500

    stelle     = request.args.get('stelle', '').strip()
    ort        = request.args.get('ort', '').strip()
    umkreis    = request.args.get('umkreis', '25')
    alter      = request.args.get('alter', '').strip()
    sortierung = request.args.get('sortierung', 'date').strip()
    page       = max(1, int(request.args.get('page', 1)))

    if not stelle:
        return jsonify({'error': 'Bitte eine Stelle eingeben.'}), 400

    app_id  = os.environ.get('ADZUNA_APP_ID') or get_setting('adzuna_app_id')
    app_key = os.environ.get('ADZUNA_APP_KEY') or get_setting('adzuna_app_key')
    if not app_id or not app_key:
        return jsonify({'error': 'Adzuna API-Key nicht konfiguriert.'}), 400

    sort_by = sortierung if sortierung in ('relevance', 'date', 'salary') else 'date'
    params = {
        'app_id':           app_id,
        'app_key':          app_key,
        'results_per_page': 50,
        'what_and':         stelle,
        'sort_by':          sort_by,
    }
    if alter:
        try:
            params['max_days_old'] = int(alter)
        except ValueError:
            pass
    if ort:
        params['where'] = ort
        try:
            params['distance'] = int(umkreis)
        except ValueError:
            params['distance'] = 25

    try:
        resp = http_requests.get(
            f'https://api.adzuna.com/v1/api/jobs/de/search/{page}',
            params=params, timeout=12,
            headers={'User-Agent': 'BewerbungsKI/1.0'}
        )
        resp.raise_for_status()
        raw = resp.json()
    except Exception as e:
        return jsonify({'error': f'Adzuna API-Fehler: {str(e)}'}), 502

    jobs_out = []
    for r in raw.get('results', []):
        created = r.get('created', '')
        age, age_days = _format_job_age(created)

        desc = r.get('description', '')
        desc = re.sub(r'<[^>]+>', ' ', desc)          # strip HTML tags
        desc = re.sub(r'\s+', ' ', desc).strip()
        full_desc = desc                              # full text for apply flow
        desc = desc[:220] + '…' if len(desc) > 220 else desc

        sal_min = r.get('salary_min')
        sal_max = r.get('salary_max')
        salary = ''
        if sal_min and sal_max:
            salary = f'{int(sal_min):,} – {int(sal_max):,} €'
        elif sal_min:
            salary = f'ab {int(sal_min):,} €'

        jobs_out.append({
            'title':        r.get('title', ''),
            'company':      r.get('company', {}).get('display_name', ''),
            'location':     r.get('location', {}).get('display_name', ''),
            'url':          r.get('redirect_url', ''),
            'age':          age,
            'age_days':     age_days,
            'description':  desc,
            'full_desc':    full_desc,
            'salary':       salary,
            'contract':     r.get('contract_type', ''),
        })

    total = raw.get('count', 0)
    suggestions = []
    if total == 0:
        import difflib
        titles_lower = {t.lower(): t for t in COMMON_JOB_TITLES}
        matches = difflib.get_close_matches(
            stelle.lower(), titles_lower.keys(), n=4, cutoff=0.4
        )
        suggestions = [titles_lower[m] for m in matches]

    return jsonify({
        'success':     True,
        'jobs':        jobs_out,
        'total':       total,
        'page':        page,
        'suggestions': suggestions,
    })




def _format_job_age(iso_str):
    if not iso_str:
        return '', 9999
    try:
        from datetime import timezone
        dt = datetime.fromisoformat(iso_str.replace('Z', '+00:00'))
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        delta = (datetime.now(timezone.utc) - dt).days
        if delta <= 0:   label = 'heute'
        elif delta == 1: label = 'gestern'
        elif delta < 7:  label = f'vor {delta} Tagen'
        elif delta < 30: label = f'vor {delta // 7} Wo.'
        else:            label = f'vor {delta // 30} Mon.'
        return label, delta
    except Exception:
        return '', 9999


# ── Bundesagentur für Arbeit API ──────────────────────────────────────────────

_BA_TOKEN_URL  = 'https://rest.arbeitsagentur.de/oauth/gettoken_cc'
_BA_JOBS_URL   = 'https://rest.arbeitsagentur.de/jobboerse/jobsuche-service/pc/v4/jobs'
_BA_DETAIL_URL = 'https://rest.arbeitsagentur.de/jobboerse/jobsuche-service/pc/v4/jobdetails/{}'

# Set BA_CLIENT_ID and BA_CLIENT_SECRET as environment variables in Railway.
# Register at: developer.arbeitsagentur.de → API "Jobsuche"
_BA_CLIENT_ID     = os.environ.get('BA_CLIENT_ID', '')
_BA_CLIENT_SECRET = os.environ.get('BA_CLIENT_SECRET', '')

_ba_token_cache = {'token': None, 'expires': 0.0}


def _get_ba_token():
    import time
    now = time.time()
    if _ba_token_cache['token'] and _ba_token_cache['expires'] > now + 60:
        return _ba_token_cache['token']
    resp = http_requests.post(
        _BA_TOKEN_URL,
        data={
            'client_id':     _BA_CLIENT_ID,
            'client_secret': _BA_CLIENT_SECRET,
            'grant_type':    'client_credentials',
        },
        timeout=10,
    )
    resp.raise_for_status()
    data = resp.json()
    token = data['access_token']
    _ba_token_cache['token']   = token
    _ba_token_cache['expires'] = now + data.get('expires_in', 3600)
    return token


@app.route('/api/jobs/search-ba')
@login_required
def jobs_search_ba():
    if not SCRAPE_SUPPORT:
        return jsonify({'error': 'requests-Bibliothek fehlt'}), 500
    if not _BA_CLIENT_ID or not _BA_CLIENT_SECRET:
        return jsonify({'error': 'BA_NOT_CONFIGURED'}), 503

    stelle     = request.args.get('stelle', '').strip()
    ort        = request.args.get('ort', '').strip()
    umkreis    = request.args.get('umkreis', '25')
    alter      = request.args.get('alter', '')
    page       = max(1, int(request.args.get('page', 1)))

    if not stelle:
        return jsonify({'error': 'Bitte eine Stelle eingeben.'}), 400

    try:
        token = _get_ba_token()
    except Exception as e:
        return jsonify({'error': f'BA API nicht erreichbar: {str(e)}'}), 502

    params = {
        'was':  stelle,
        'size': 25,
        'page': page - 1,   # BA uses 0-based pages
    }
    if ort:
        params['wo'] = ort
        try:
            params['umkreis'] = int(umkreis)
        except ValueError:
            params['umkreis'] = 25
    if alter:
        try:
            params['veroeffentlichtseit'] = int(alter)
        except ValueError:
            pass

    headers = {
        'Authorization': f'Bearer {token}',
        'User-Agent':    'BewerbungsKI/1.0',
    }

    try:
        resp = http_requests.get(_BA_JOBS_URL, params=params, headers=headers, timeout=12)
        resp.raise_for_status()
        raw = resp.json()
    except Exception as e:
        return jsonify({'error': f'BA API Fehler: {str(e)}'}), 502

    jobs_out = []
    for r in raw.get('stellenangebote', []) or []:
        created  = r.get('aktuelleVeroeffentlichungsdatum', '')
        age, age_days = _format_job_age(created)

        arbeitsort = r.get('arbeitsort', {}) or {}
        loc_parts  = [arbeitsort.get('ort', ''), arbeitsort.get('region', '')]
        location   = ', '.join(p for p in loc_parts if p)

        modelle  = r.get('arbeitszeitmodelle', []) or []
        contract = ', '.join(modelle) if modelle else ''

        hash_id = r.get('hashId', '')
        ref_nr  = r.get('refnr', '')

        jobs_out.append({
            'title':    r.get('titel', ''),
            'company':  r.get('arbeitgeber', ''),
            'location': location,
            'url':      f'https://www.arbeitsagentur.de/jobsuche/stelle/{hash_id}' if hash_id else '',
            'hash_id':  hash_id,
            'ref_nr':   ref_nr,
            'age':      age,
            'age_days': age_days,
            'description': '',
            'full_desc':   '',
            'salary':   '',
            'contract': contract,
            'source':   'ba',
        })

    total = raw.get('maxErgebnisse', 0)

    suggestions = []
    if total == 0:
        import difflib
        titles_lower = {t.lower(): t for t in COMMON_JOB_TITLES}
        matches = difflib.get_close_matches(stelle.lower(), titles_lower.keys(), n=4, cutoff=0.4)
        suggestions = [titles_lower[m] for m in matches]

    return jsonify({
        'success':     True,
        'jobs':        jobs_out,
        'total':       total,
        'page':        page,
        'suggestions': suggestions,
        'source':      'ba',
    })


@app.route('/api/jobs/detail-ba/<hash_id>')
@login_required
def job_detail_ba(hash_id):
    if not SCRAPE_SUPPORT:
        return jsonify({'error': 'requests-Bibliothek fehlt'}), 500
    if not _BA_CLIENT_ID or not _BA_CLIENT_SECRET:
        return jsonify({'error': 'BA_NOT_CONFIGURED'}), 503

    try:
        token = _get_ba_token()
    except Exception as e:
        return jsonify({'error': f'BA API nicht erreichbar: {str(e)}'}), 502

    headers = {
        'Authorization': f'Bearer {token}',
        'User-Agent':    'BewerbungsKI/1.0',
    }

    try:
        resp = http_requests.get(
            _BA_DETAIL_URL.format(hash_id),
            headers=headers, timeout=12,
        )
        resp.raise_for_status()
        data = resp.json()
    except Exception as e:
        return jsonify({'error': f'Detail konnte nicht geladen werden: {str(e)}'}), 502

    raw_desc = data.get('stellenbeschreibung', '') or ''
    # Strip HTML tags if present
    clean = re.sub(r'<[^>]+>', '\n', raw_desc)
    clean = re.sub(r'&amp;',  '&',  clean)
    clean = re.sub(r'&lt;',   '<',  clean)
    clean = re.sub(r'&gt;',   '>',  clean)
    clean = re.sub(r'&nbsp;', ' ',  clean)
    clean = re.sub(r'\n{3,}', '\n\n', clean).strip()

    # Enrich with salary if available
    verguetung = data.get('verguetung', '') or ''

    return jsonify({
        'success':     True,
        'description': clean,
        'salary':      verguetung,
    })


# ── API: Anschreiben generieren ───────────────────────────────────────────────

@app.route('/api/generate', methods=['POST'])
@login_required
def generate():
    data = request.get_json()
    stellenanzeige = (data.get('stellenanzeige') or '').strip()
    firma = (data.get('firma') or '').strip()
    stelle = (data.get('stelle') or '').strip()

    if not stellenanzeige:
        return jsonify({'error': 'Bitte eine Stellenanzeige einfügen.'}), 400

    api_key = os.environ.get('ANTHROPIC_API_KEY')
    if not api_key:
        return jsonify({'error': 'Anthropic API-Key nicht konfiguriert (Server-Fehler).'}), 500

    cv_text = get_setting('cv_text')
    if not cv_text:
        return jsonify({'error': 'Kein Lebenslauf hochgeladen. Bitte zuerst hochladen.'}), 400

    name = get_setting('name', 'Bewerberin/Bewerber')
    ort  = get_setting('ort', 'Meine Stadt')

    try:
        client = anthropic.Anthropic(api_key=api_key)
        prompt = f"""Du bist ein Experte für professionelle deutsche Bewerbungsschreiben.

Erstelle ein vollständiges, überzeugendes Anschreiben basierend auf:

**LEBENSLAUF:**
{cv_text[:3500]}

**STELLENANZEIGE ({firma} – {stelle}):**
{stellenanzeige[:2500]}

**BEWERBER:** {name}, wohnhaft in {ort}

Anforderungen:
- Professionell, persönlich und überzeugend
- Relevante Qualifikationen aus dem Lebenslauf gezielt hervorheben
- Konkrete Bezüge zur Stellenanzeige herstellen
- Deutsche Bewerbungsstandards einhalten
- Max. 180 Wörter, kurz und prägnant, passt auf eine Seite
- Mit passender Anrede beginnen
- Mit "Mit freundlichen Grüßen," + Leerzeile + Name enden
- Keine Gedankenstriche (–) oder Bindestriche (-) als Satzzeichen verwenden; stattdessen Komma oder Punkt nutzen

Gib NUR den Brieftext zurück (von Anrede bis Unterschrift), ohne Absenderblock oder Datum."""

        msg = client.messages.create(
            model='claude-sonnet-4-6',
            max_tokens=1024,
            messages=[{'role': 'user', 'content': prompt}]
        )
        return jsonify({'success': True, 'anschreiben': msg.content[0].text})

    except anthropic.AuthenticationError:
        return jsonify({'error': 'Ungültiger API-Key. Bitte in den Einstellungen prüfen.'}), 400
    except Exception as e:
        return jsonify({'error': f'Fehler: {str(e)}'}), 500


# ── API: AI Chat ──────────────────────────────────────────────────────────────

@app.route('/api/chat/<int:bid>', methods=['POST'])
@login_required
def chat(bid):
    data = request.get_json()
    message = (data.get('message') or '').strip()

    if not message:
        return jsonify({'error': 'Nachricht fehlt'}), 400

    api_key = os.environ.get('ANTHROPIC_API_KEY')
    if not api_key:
        return jsonify({'error': 'Anthropic API-Key nicht konfiguriert (Server-Fehler).'}), 500

    conn = get_db()
    b = conn.execute('SELECT * FROM bewerbungen WHERE id = ? AND user_id = ?', (bid, current_user.id)).fetchone()
    conn.close()

    if not b:
        return jsonify({'error': 'Bewerbung nicht gefunden'}), 404

    cv_text  = get_setting('cv_text', '')
    anschreiben = b['anschreiben'] or ''

    system = f"""Du bist ein Experte für deutsche Bewerbungsschreiben und hilfst dem Nutzer,
sein Anschreiben zu verbessern.

Kontext:
- Stelle: {b['stelle']} bei {b['firma']}
- Stellenanzeige (Auszug): {(b['stellenanzeige'] or '')[:800]}
- Lebenslauf (Auszug): {cv_text[:1200]}

Aktuelles Anschreiben:
---
{anschreiben}
---

Stilregel: Keine Gedankenstriche (–) oder Bindestriche (-) als Satzzeichen; stattdessen Komma oder Punkt.
Längenvorgabe: Max. 180 Wörter, sofern der Nutzer nicht ausdrücklich mehr verlangt.

Wenn der Nutzer eine Änderung oder Überarbeitung wünscht:
1. Erkläre in 1-2 Sätzen was du geändert hast.
2. Gib das vollständig überarbeitete Anschreiben aus, genau so formatiert:
[ANSCHREIBEN_START]
<vollständiges Anschreiben hier>
[ANSCHREIBEN_ENDE]

Wenn der Nutzer eine Frage stellt oder nur Feedback gibt, antworte kurz ohne die Tags."""

    try:
        client = anthropic.Anthropic(api_key=api_key)
        msg = client.messages.create(
            model='claude-sonnet-4-6',
            max_tokens=1500,
            system=system,
            messages=[{'role': 'user', 'content': message}]
        )
        response = msg.content[0].text

        updated_anschreiben = None
        display_text = response

        if '[ANSCHREIBEN_START]' in response and '[ANSCHREIBEN_ENDE]' in response:
            s = response.index('[ANSCHREIBEN_START]') + len('[ANSCHREIBEN_START]')
            e = response.index('[ANSCHREIBEN_ENDE]')
            updated_anschreiben = response[s:e].strip()

            display_text = response[:response.index('[ANSCHREIBEN_START]')].strip()
            if not display_text:
                display_text = 'Anschreiben wurde aktualisiert.'

            conn = get_db()
            conn.execute('UPDATE bewerbungen SET anschreiben = ? WHERE id = ?',
                         (updated_anschreiben, bid))
            conn.commit()
            conn.close()

        # Persist chat messages
        conn = get_db()
        conn.execute(
            'INSERT INTO chat_nachrichten (bewerbung_id, rolle, inhalt) VALUES (?,?,?)',
            (bid, 'user', message)
        )
        conn.execute(
            'INSERT INTO chat_nachrichten (bewerbung_id, rolle, inhalt) VALUES (?,?,?)',
            (bid, 'assistant', display_text)
        )
        conn.commit()
        conn.close()

        return jsonify({
            'success': True,
            'message': display_text,
            'updated_anschreiben': updated_anschreiben,
        })

    except anthropic.AuthenticationError:
        return jsonify({'error': 'Ungültiger API-Key'}), 400
    except Exception as e:
        return jsonify({'error': str(e)}), 500


# ── API: Bewerbungen CRUD ─────────────────────────────────────────────────────

@app.route('/api/save', methods=['POST'])
@login_required
def save():
    data = request.get_json()
    uid = current_user.id
    conn = get_db()
    params = (uid, data.get('firma', ''), data.get('stelle', ''), data.get('url', ''),
              data.get('stellenanzeige', ''), data.get('anschreiben', ''),
              data.get('status', 'Entwurf'), data.get('notizen', ''),
              data.get('datum_gesendet') or None)
    if _PG:
        cur = conn.execute(
            'INSERT INTO bewerbungen (user_id, firma, stelle, url, stellenanzeige, anschreiben, status, notizen, datum_gesendet)'
            ' VALUES (?,?,?,?,?,?,?,?,?) RETURNING id', params
        )
        bid = cur.fetchone()['id']
    else:
        cur = conn.execute(
            'INSERT INTO bewerbungen (user_id, firma, stelle, url, stellenanzeige, anschreiben, status, notizen, datum_gesendet)'
            ' VALUES (?,?,?,?,?,?,?,?,?)', params
        )
        bid = cur.lastrowid
    conn.commit()
    conn.close()
    return jsonify({'success': True, 'id': bid})


@app.route('/api/update/<int:bid>', methods=['POST'])
@login_required
def update(bid):
    data = request.get_json()
    conn = get_db()
    conn.execute(
        '''UPDATE bewerbungen SET firma=?, stelle=?, url=?, stellenanzeige=?,
           anschreiben=?, status=?, notizen=?, datum_gesendet=?
           WHERE id=? AND user_id=?''',
        (data.get('firma', ''), data.get('stelle', ''), data.get('url', ''),
         data.get('stellenanzeige', ''), data.get('anschreiben', ''),
         data.get('status', 'Entwurf'), data.get('notizen', ''),
         data.get('datum_gesendet') or None, bid, current_user.id)
    )
    conn.commit()
    conn.close()
    return jsonify({'success': True})


@app.route('/api/update-status/<int:bid>', methods=['POST'])
@login_required
def update_status(bid):
    data = request.get_json()
    status = data.get('status', '')
    if status not in STATUS_OPTIONS:
        return jsonify({'error': 'Ungültiger Status'}), 400
    conn = get_db()
    conn.execute('UPDATE bewerbungen SET status = ? WHERE id = ? AND user_id = ?',
                 (status, bid, current_user.id))
    conn.commit()
    conn.close()
    return jsonify({'success': True})


@app.route('/api/delete/<int:bid>', methods=['POST'])
@login_required
def delete(bid):
    conn = get_db()
    conn.execute('DELETE FROM bewerbungen WHERE id = ? AND user_id = ?', (bid, current_user.id))
    conn.commit()
    conn.close()
    return jsonify({'success': True})


# ── CV Builder ────────────────────────────────────────────────────────────────

_CV_EMPTY = {
    'personal': {'name': '', 'email': '', 'phone': '', 'address': '', 'city': '', 'linkedin': '', 'website': ''},
    'summary': '',
    'experience': [],
    'education': [],
    'skills': [],
    'languages': [],
    'certifications': [],
}


def _get_cv_json():
    if not current_user.is_authenticated:
        return {}, 'klassisch'
    conn = get_db()
    row = conn.execute('SELECT data, template FROM cv_data WHERE user_id = ?', (current_user.id,)).fetchone()
    conn.close()
    if row:
        try:
            return json.loads(row['data'] or '{}'), row['template'] or 'klassisch'
        except Exception:
            pass
    return {}, 'klassisch'


def _save_cv_json(data, template):
    if not current_user.is_authenticated:
        return
    uid = current_user.id
    payload = (uid, json.dumps(data, ensure_ascii=False), template)
    conn = get_db()
    if _PG:
        conn.execute(
            'INSERT INTO cv_data (user_id, data, template, updated_at) VALUES (?, ?, ?, CURRENT_TIMESTAMP::TEXT)'
            ' ON CONFLICT (user_id) DO UPDATE SET data = EXCLUDED.data,'
            ' template = EXCLUDED.template, updated_at = CURRENT_TIMESTAMP::TEXT',
            payload
        )
    else:
        conn.execute(
            'INSERT OR REPLACE INTO cv_data (user_id, data, template, updated_at) VALUES (?, ?, ?, datetime("now"))',
            payload
        )
    conn.commit()
    conn.close()


@app.route('/api/parse-cv', methods=['POST'])
@login_required
def parse_cv():
    api_key = os.environ.get('ANTHROPIC_API_KEY')
    if not api_key:
        return jsonify({'error': 'Anthropic API-Key nicht konfiguriert (Server-Fehler).'}), 500

    cv_text = get_setting('cv_text')
    if not cv_text:
        return jsonify({'error': 'Kein Lebenslauf hochgeladen. Bitte zuerst hochladen.'}), 400

    prompt = f"""Analysiere diesen Lebenslauf und extrahiere ALLE Informationen strukturiert.

Lebenslauf-Text:
{cv_text[:5000]}

Antworte NUR mit validem JSON in diesem exakten Format (keine Erklärungen, kein Markdown, nur reines JSON):
{{
  "personal": {{"name": "", "email": "", "phone": "", "address": "", "city": "", "linkedin": "", "website": ""}},
  "summary": "",
  "experience": [
    {{
      "company": "", "role": "", "location": "",
      "start": "MM/YYYY", "end": "MM/YYYY", "current": false,
      "bullets": ["Aufgabe 1", "Aufgabe 2"]
    }}
  ],
  "education": [
    {{"institution": "", "degree": "", "field": "", "start": "YYYY", "end": "YYYY", "grade": ""}}
  ],
  "skills": ["Skill 1", "Skill 2"],
  "languages": [{{"lang": "", "level": ""}}],
  "certifications": [{{"name": "", "issuer": "", "year": ""}}]
}}

Regeln:
- Nutze leere Strings "" oder [] wenn etwas nicht vorhanden
- current=true wenn noch tätig
- Gib NUR das JSON zurück, sonst nichts"""

    try:
        client = anthropic.Anthropic(api_key=api_key)
        msg = client.messages.create(
            model='claude-sonnet-4-6',
            max_tokens=2048,
            messages=[{'role': 'user', 'content': prompt}]
        )
        raw = msg.content[0].text.strip()
        # Strip markdown code fences if present
        if raw.startswith('```'):
            raw = re.sub(r'^```[a-z]*\n?', '', raw)
            raw = re.sub(r'\n?```$', '', raw)

        parsed = json.loads(raw)
        _save_cv_json(parsed, 'klassisch')
        return jsonify({'success': True, 'data': parsed})

    except json.JSONDecodeError as e:
        return jsonify({'error': f'KI hat ungültiges JSON zurückgegeben: {str(e)}'}), 500
    except anthropic.AuthenticationError:
        return jsonify({'error': 'Ungültiger API-Key.'}), 400
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/api/cv-data', methods=['GET', 'POST'])
@login_required
def cv_data_api():
    if request.method == 'GET':
        data, template = _get_cv_json()
        return jsonify({'success': True, 'data': data, 'template': template})

    body = request.get_json()
    data = body.get('data', {})
    template = body.get('template', 'klassisch')
    if template not in ('klassisch', 'modern', 'minimal'):
        template = 'klassisch'
    _save_cv_json(data, template)
    return jsonify({'success': True})


@app.route('/cv/pdf/<template>')
@login_required
def cv_pdf(template):
    if template not in ('klassisch', 'modern', 'minimal'):
        return 'Ungültige Vorlage', 400
    data, _ = _get_cv_json()
    if not data:
        return 'Kein Lebenslauf gespeichert', 400

    s = {k: get_setting(k) for k in ['name', 'email', 'telefon', 'strasse', 'plz_ort', 'ort']}

    builders = {
        'klassisch': _cv_pdf_klassisch,
        'modern':    _cv_pdf_modern,
        'minimal':   _cv_pdf_minimal,
    }
    buf = builders[template](data, s)
    name = (data.get('personal', {}).get('name') or s.get('name') or 'Lebenslauf').replace(' ', '_')
    fname = f'Lebenslauf_{name}_{template}.pdf'
    return send_file(buf, as_attachment=True, download_name=fname, mimetype='application/pdf')


# ── CV PDF Templates ──────────────────────────────────────────────────────────

NAVY  = colors.HexColor('#1e3a5f')
CYAN  = colors.HexColor('#06b6d4')
SLATE = colors.HexColor('#64748b')
INK   = colors.HexColor('#0f172a')
LIGHT = colors.HexColor('#f8fafc')
RULE  = colors.HexColor('#e2e8f0')


def _cv_pdf_klassisch(data, profile):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=2.5*cm, rightMargin=2.5*cm,
                            topMargin=2*cm, bottomMargin=2*cm)
    p = data.get('personal', {})

    name_s    = ParagraphStyle('K_name', fontName='Helvetica-Bold', fontSize=22,
                                alignment=1, spaceAfter=4, textColor=INK)
    contact_s = ParagraphStyle('K_contact', fontName='Helvetica', fontSize=10,
                                alignment=1, textColor=SLATE, spaceAfter=14)
    section_s = ParagraphStyle('K_sec', fontName='Helvetica-Bold', fontSize=10,
                                textColor=NAVY, spaceBefore=14, spaceAfter=6)
    body_s    = ParagraphStyle('K_body', fontName='Helvetica', fontSize=10, leading=14, spaceAfter=2)
    bold_s    = ParagraphStyle('K_bold', fontName='Helvetica-Bold', fontSize=10, leading=14, spaceAfter=2)
    small_s   = ParagraphStyle('K_small', fontName='Helvetica', fontSize=9.5,
                                textColor=SLATE, leading=13, spaceAfter=2)
    bullet_s  = ParagraphStyle('K_bul', fontName='Helvetica', fontSize=10, leading=14,
                                leftIndent=10, spaceAfter=1)
    summary_s = ParagraphStyle('K_sum', fontName='Helvetica', fontSize=10.5, leading=16, spaceAfter=12)

    story = []

    name = p.get('name') or profile.get('name') or ''
    if name:
        story.append(Paragraph(name, name_s))

    contacts = [p.get('phone') or profile.get('telefon'), p.get('email') or profile.get('email'),
                p.get('city') or profile.get('plz_ort'), p.get('linkedin')]
    ct = ' · '.join(x for x in contacts if x)
    if ct:
        story.append(Paragraph(ct, contact_s))

    story.append(HRFlowable(width='100%', thickness=1.2, color=NAVY, spaceAfter=10))

    if data.get('summary'):
        story.append(Paragraph(data['summary'], summary_s))
        story.append(HRFlowable(width='100%', thickness=0.5, color=RULE, spaceAfter=4))

    def two_col(date_txt, parts):
        tbl = Table([[Paragraph(date_txt, small_s), parts]], colWidths=[3*cm, None])
        tbl.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('TOPPADDING', (0, 0), (-1, -1), 0),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
            ('LEFTPADDING', (0, 0), (-1, -1), 0),
            ('RIGHTPADDING', (0, 0), (0, -1), 12),
            ('RIGHTPADDING', (1, 0), (1, -1), 0),
        ]))
        return tbl

    if data.get('experience'):
        story.append(Paragraph('BERUFSERFAHRUNG', section_s))
        story.append(HRFlowable(width='100%', thickness=0.5, color=NAVY, spaceAfter=8))
        for exp in data['experience']:
            end = 'heute' if exp.get('current') else (exp.get('end') or '')
            date_txt = f"{exp.get('start', '')} –\n{end}"
            parts = []
            if exp.get('role'):
                parts.append(Paragraph(f"<b>{exp['role']}</b>", bold_s))
            cl = ' · '.join(filter(None, [exp.get('company'), exp.get('location')]))
            if cl:
                parts.append(Paragraph(cl, small_s))
            for b in (exp.get('bullets') or []):
                if b.strip():
                    parts.append(Paragraph(f'• {b}', bullet_s))
            if parts:
                story.append(two_col(date_txt, parts))
                story.append(Spacer(1, 8))

    if data.get('education'):
        story.append(Paragraph('AUSBILDUNG', section_s))
        story.append(HRFlowable(width='100%', thickness=0.5, color=NAVY, spaceAfter=8))
        for edu in data['education']:
            date_txt = f"{edu.get('start', '')} –\n{edu.get('end', '')}"
            parts = []
            df = ' – '.join(filter(None, [edu.get('degree'), edu.get('field')]))
            if df:
                parts.append(Paragraph(f'<b>{df}</b>', bold_s))
            inst = edu.get('institution', '')
            if edu.get('grade'):
                inst += f' · Note: {edu["grade"]}'
            if inst:
                parts.append(Paragraph(inst, small_s))
            if parts:
                story.append(two_col(date_txt, parts))
                story.append(Spacer(1, 8))

    if data.get('skills'):
        story.append(Paragraph('KENNTNISSE', section_s))
        story.append(HRFlowable(width='100%', thickness=0.5, color=NAVY, spaceAfter=8))
        story.append(Paragraph(' · '.join(s for s in data['skills'] if s), body_s))

    if data.get('languages'):
        story.append(Paragraph('SPRACHEN', section_s))
        story.append(HRFlowable(width='100%', thickness=0.5, color=NAVY, spaceAfter=8))
        lt = ' · '.join(f"{l.get('lang', '')}: {l.get('level', '')}" for l in data['languages'] if l.get('lang'))
        story.append(Paragraph(lt, body_s))

    if data.get('certifications'):
        story.append(Paragraph('ZERTIFIKATE', section_s))
        story.append(HRFlowable(width='100%', thickness=0.5, color=NAVY, spaceAfter=8))
        for cert in data['certifications']:
            txt = cert.get('name', '')
            if cert.get('issuer'):
                txt += f' · {cert["issuer"]}'
            if cert.get('year'):
                txt += f' ({cert["year"]})'
            story.append(Paragraph(txt, body_s))

    doc.build(story)
    buf.seek(0)
    return buf


def _cv_pdf_modern(data, profile):
    buf = io.BytesIO()
    L, R = 2.5*cm, 2.5*cm
    HDR_H = 3.6*cm

    # topMargin reserves space for the drawn header; header is painted via canvas callback
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=L, rightMargin=R,
                            topMargin=HDR_H + 0.4*cm, bottomMargin=1.5*cm)
    p = data.get('personal', {})
    CONTENT_W = A4[0] - L - R

    name = p.get('name') or profile.get('name') or ''
    contacts = [p.get('phone') or profile.get('telefon'), p.get('email') or profile.get('email'),
                p.get('city') or profile.get('plz_ort'), p.get('linkedin')]
    ct = '  ·  '.join(x for x in contacts if x)

    def draw_header(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(NAVY)
        canvas.rect(0, A4[1] - HDR_H, A4[0], HDR_H, fill=1, stroke=0)
        y_name = A4[1] - 1.7*cm
        if name:
            canvas.setFillColor(colors.white)
            canvas.setFont('Helvetica-Bold', 22)
            canvas.drawString(L, y_name, name)
        if ct:
            canvas.setFillColor(colors.HexColor('#94a3b8'))
            canvas.setFont('Helvetica', 10)
            canvas.drawString(L, y_name - 0.8*cm, ct[:120])
        canvas.restoreState()

    sec_s    = ParagraphStyle('M_sec', fontName='Helvetica-Bold', fontSize=10,
                               textColor=NAVY, spaceBefore=14, spaceAfter=6)
    body_s   = ParagraphStyle('M_body', fontName='Helvetica', fontSize=10, leading=14, spaceAfter=2)
    bold_s   = ParagraphStyle('M_bold', fontName='Helvetica-Bold', fontSize=10, leading=14, spaceAfter=2)
    small_s  = ParagraphStyle('M_small', fontName='Helvetica', fontSize=9.5, textColor=SLATE, leading=13)
    bullet_s = ParagraphStyle('M_bul', fontName='Helvetica', fontSize=10, leading=14,
                               leftIndent=10, spaceAfter=1)
    sum_s    = ParagraphStyle('M_sum', fontName='Helvetica', fontSize=10.5, leading=16, spaceAfter=12)

    def section_title(txt):
        bar = Table([[Paragraph(txt.upper(), sec_s)]], colWidths=[CONTENT_W])
        bar.setStyle(TableStyle([
            ('LINEBEFORE', (0, 0), (0, -1), 3, CYAN),
            ('LEFTPADDING', (0, 0), (-1, -1), 8),
            ('TOPPADDING', (0, 0), (-1, -1), 2),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
        ]))
        return bar

    def two_col(date_txt, parts):
        tbl = Table([[Paragraph(date_txt, small_s), parts]], colWidths=[3*cm, None])
        tbl.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('TOPPADDING', (0, 0), (-1, -1), 0),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
            ('LEFTPADDING', (0, 0), (-1, -1), 0),
            ('RIGHTPADDING', (0, 0), (0, -1), 12),
            ('RIGHTPADDING', (1, 0), (1, -1), 0),
        ]))
        return tbl

    story = []

    if data.get('summary'):
        story.append(Paragraph(data['summary'], sum_s))

    if data.get('experience'):
        story.append(section_title('Berufserfahrung'))
        story.append(HRFlowable(width='100%', thickness=0.5, color=RULE, spaceAfter=8))
        for exp in data['experience']:
            end = 'heute' if exp.get('current') else (exp.get('end') or '')
            date_txt = f"{exp.get('start', '')} –\n{end}"
            parts = []
            if exp.get('role'):
                parts.append(Paragraph(f"<b>{exp['role']}</b>", bold_s))
            cl = ' · '.join(filter(None, [exp.get('company'), exp.get('location')]))
            if cl:
                parts.append(Paragraph(cl, small_s))
            for b in (exp.get('bullets') or []):
                if b.strip():
                    parts.append(Paragraph(f'• {b}', bullet_s))
            if parts:
                story.append(two_col(date_txt, parts))
                story.append(Spacer(1, 8))

    if data.get('education'):
        story.append(section_title('Ausbildung'))
        story.append(HRFlowable(width='100%', thickness=0.5, color=RULE, spaceAfter=8))
        for edu in data['education']:
            date_txt = f"{edu.get('start', '')} –\n{edu.get('end', '')}"
            parts = []
            df = ' – '.join(filter(None, [edu.get('degree'), edu.get('field')]))
            if df:
                parts.append(Paragraph(f'<b>{df}</b>', bold_s))
            inst = edu.get('institution', '')
            if edu.get('grade'):
                inst += f' · Note: {edu["grade"]}'
            if inst:
                parts.append(Paragraph(inst, small_s))
            if parts:
                story.append(two_col(date_txt, parts))
                story.append(Spacer(1, 8))

    if data.get('skills'):
        story.append(section_title('Kenntnisse'))
        story.append(HRFlowable(width='100%', thickness=0.5, color=RULE, spaceAfter=8))
        story.append(Paragraph(' · '.join(s for s in data['skills'] if s), body_s))
        story.append(Spacer(1, 4))

    if data.get('languages'):
        story.append(section_title('Sprachen'))
        story.append(HRFlowable(width='100%', thickness=0.5, color=RULE, spaceAfter=8))
        lt = ' · '.join(f"{l.get('lang', '')}: {l.get('level', '')}" for l in data['languages'] if l.get('lang'))
        story.append(Paragraph(lt, body_s))

    if data.get('certifications'):
        story.append(section_title('Zertifikate'))
        story.append(HRFlowable(width='100%', thickness=0.5, color=RULE, spaceAfter=8))
        for cert in data['certifications']:
            txt = cert.get('name', '')
            if cert.get('issuer'):
                txt += f' · {cert["issuer"]}'
            if cert.get('year'):
                txt += f' ({cert["year"]})'
            story.append(Paragraph(txt, body_s))

    doc.build(story, onFirstPage=draw_header, onLaterPages=draw_header)
    buf.seek(0)
    return buf


def _cv_pdf_minimal(data, profile):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=3*cm, rightMargin=3*cm,
                            topMargin=2.5*cm, bottomMargin=2*cm)
    p = data.get('personal', {})

    name_s   = ParagraphStyle('N_name', fontName='Helvetica-Bold', fontSize=26,
                               textColor=INK, spaceAfter=6)
    ct_s     = ParagraphStyle('N_ct', fontName='Helvetica', fontSize=10,
                               textColor=SLATE, spaceAfter=20)
    sec_s    = ParagraphStyle('N_sec', fontName='Helvetica-Bold', fontSize=10.5,
                               textColor=NAVY, spaceBefore=18, spaceAfter=8)
    body_s   = ParagraphStyle('N_body', fontName='Helvetica', fontSize=10.5, leading=16, spaceAfter=2)
    bold_s   = ParagraphStyle('N_bold', fontName='Helvetica-Bold', fontSize=10.5, leading=16, spaceAfter=2)
    small_s  = ParagraphStyle('N_small', fontName='Helvetica', fontSize=10, textColor=SLATE, leading=14)
    bullet_s = ParagraphStyle('N_bul', fontName='Helvetica', fontSize=10.5, leading=16,
                               leftIndent=12, spaceAfter=2)
    sum_s    = ParagraphStyle('N_sum', fontName='Helvetica', fontSize=11, leading=18,
                               spaceAfter=16, textColor=colors.HexColor('#334155'))

    story = []

    # Cyan top accent
    story.append(HRFlowable(width='100%', thickness=3, color=CYAN, spaceAfter=20))

    name = p.get('name') or profile.get('name') or ''
    if name:
        story.append(Paragraph(name, name_s))

    contacts = [p.get('phone') or profile.get('telefon'), p.get('email') or profile.get('email'),
                p.get('city') or profile.get('plz_ort'), p.get('linkedin')]
    ct = '  ·  '.join(x for x in contacts if x)
    if ct:
        story.append(Paragraph(ct, ct_s))

    story.append(HRFlowable(width='100%', thickness=0.5, color=RULE, spaceAfter=8))

    if data.get('summary'):
        story.append(Paragraph(data['summary'], sum_s))

    CONTENT_W = A4[0] - 6*cm

    def two_col(date_txt, parts):
        tbl = Table([[Paragraph(date_txt, small_s), parts]], colWidths=[3.2*cm, None])
        tbl.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('TOPPADDING', (0, 0), (-1, -1), 0),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
            ('LEFTPADDING', (0, 0), (-1, -1), 0),
            ('RIGHTPADDING', (0, 0), (0, -1), 16),
            ('RIGHTPADDING', (1, 0), (1, -1), 0),
        ]))
        return tbl

    if data.get('experience'):
        story.append(Paragraph('Berufserfahrung', sec_s))
        story.append(HRFlowable(width='100%', thickness=0.5, color=RULE, spaceAfter=10))
        for exp in data['experience']:
            end = 'heute' if exp.get('current') else (exp.get('end') or '')
            date_txt = f"{exp.get('start', '')} – {end}"
            parts = []
            if exp.get('role'):
                parts.append(Paragraph(f"<b>{exp['role']}</b>", bold_s))
            cl = ' · '.join(filter(None, [exp.get('company'), exp.get('location')]))
            if cl:
                parts.append(Paragraph(cl, small_s))
            for b in (exp.get('bullets') or []):
                if b.strip():
                    parts.append(Paragraph(f'• {b}', bullet_s))
            if parts:
                story.append(two_col(date_txt, parts))
                story.append(Spacer(1, 10))

    if data.get('education'):
        story.append(Paragraph('Ausbildung', sec_s))
        story.append(HRFlowable(width='100%', thickness=0.5, color=RULE, spaceAfter=10))
        for edu in data['education']:
            date_txt = f"{edu.get('start', '')} – {edu.get('end', '')}"
            parts = []
            df = ' – '.join(filter(None, [edu.get('degree'), edu.get('field')]))
            if df:
                parts.append(Paragraph(f'<b>{df}</b>', bold_s))
            inst = edu.get('institution', '')
            if edu.get('grade'):
                inst += f' · Note: {edu["grade"]}'
            if inst:
                parts.append(Paragraph(inst, small_s))
            if parts:
                story.append(two_col(date_txt, parts))
                story.append(Spacer(1, 10))

    if data.get('skills'):
        story.append(Paragraph('Kenntnisse', sec_s))
        story.append(HRFlowable(width='100%', thickness=0.5, color=RULE, spaceAfter=10))
        story.append(Paragraph(' · '.join(s for s in data['skills'] if s), body_s))

    if data.get('languages'):
        story.append(Paragraph('Sprachen', sec_s))
        story.append(HRFlowable(width='100%', thickness=0.5, color=RULE, spaceAfter=10))
        lt = ' · '.join(f"{l.get('lang', '')}: {l.get('level', '')}" for l in data['languages'] if l.get('lang'))
        story.append(Paragraph(lt, body_s))

    if data.get('certifications'):
        story.append(Paragraph('Zertifikate', sec_s))
        story.append(HRFlowable(width='100%', thickness=0.5, color=RULE, spaceAfter=10))
        for cert in data['certifications']:
            txt = cert.get('name', '')
            if cert.get('issuer'):
                txt += f' · {cert["issuer"]}'
            if cert.get('year'):
                txt += f' ({cert["year"]})'
            story.append(Paragraph(txt, body_s))

    doc.build(story)
    buf.seek(0)
    return buf


# ── Anschreiben PDF Export ─────────────────────────────────────────────────────

@app.route('/bewerbung/<int:bid>/pdf')
@login_required
def export_pdf(bid):
    conn = get_db()
    row = conn.execute('SELECT * FROM bewerbungen WHERE id = ? AND user_id = ?', (bid, current_user.id)).fetchone()
    conn.close()
    if not row:
        return 'Nicht gefunden', 404

    b = dict(row)
    s = {k: get_setting(k) for k in ['name', 'email', 'telefon', 'strasse', 'plz_ort', 'ort']}
    buf = _build_pdf(b, s)

    safe = f"Anschreiben_{b['firma']}_{b['stelle']}".replace(' ', '_').replace('/', '-')
    return send_file(buf, as_attachment=True, download_name=f'{safe}.pdf',
                     mimetype='application/pdf')


def _build_pdf(b, s):
    buf = io.BytesIO()
    # DIN 5008: left=2.5cm, right=2cm, top=2.7cm, bottom=2cm
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            rightMargin=2*cm, leftMargin=2.5*cm,
                            topMargin=2.7*cm, bottomMargin=2*cm)

    ACCENT     = colors.HexColor('#1a56db')
    INK        = colors.HexColor('#111827')
    GREY       = colors.HexColor('#6b7280')
    ACCENT_HEX = '1a56db'
    GREY_HEX   = '6b7280'

    body_w = A4[0] - 2.5*cm - 2*cm  # usable width ~165mm

    def _style(name, **kw):
        base = ParagraphStyle(name, fontName='Helvetica', fontSize=11,
                              leading=16, textColor=INK)
        for k, v in kw.items():
            setattr(base, k, v)
        return base

    P     = _style('P')
    PBold = _style('PB', fontName='Helvetica-Bold')
    PSmall= _style('PS', fontSize=9, leading=13, textColor=GREY)
    PRight= _style('PR', alignment=2)
    PSubj = _style('PSub', fontName='Helvetica-Bold', fontSize=12, leading=18, textColor=INK)

    story = []

    # ── Header: two-column (sender left | contact right) ──────────────────────
    sender_lines = []
    if s.get('name'):
        sender_lines.append(f"<font name='Helvetica-Bold' size='13' color='#{ACCENT_HEX}'>{s['name']}</font>")
    for field in ['strasse', 'plz_ort']:
        if s.get(field):
            sender_lines.append(s[field])

    contact_lines = []
    if s.get('telefon'):
        contact_lines.append(f"<font color='#{GREY_HEX}'>Tel:</font>  {s['telefon']}")
    if s.get('email'):
        contact_lines.append(f"<font color='#{GREY_HEX}'>Mail:</font> {s['email']}")

    if sender_lines or contact_lines:
        left_cell  = [Paragraph(l, PSmall) for l in sender_lines] or [Paragraph('', PSmall)]
        right_cell = [Paragraph(l, PSmall) for l in contact_lines] or [Paragraph('', PSmall)]
        hdr_table = Table([[left_cell, right_cell]], colWidths=[body_w*0.6, body_w*0.4])
        hdr_table.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
            ('LEFTPADDING',  (0,0), (-1,-1), 0),
            ('RIGHTPADDING', (0,0), (-1,-1), 0),
            ('BOTTOMPADDING',(0,0), (-1,-1), 0),
            ('TOPPADDING',   (0,0), (-1,-1), 0),
        ]))
        story.append(hdr_table)
        story.append(Spacer(1, 0.25*cm))
        story.append(HRFlowable(width='100%', thickness=1.5, color=ACCENT, spaceAfter=0.5*cm))

    # ── Recipient block (DIN 5008 address window zone) ─────────────────────────
    recipient_lines = []
    if b.get('firma'):
        recipient_lines.append(Paragraph(b['firma'], PBold))
    # placeholder for address if we ever store it
    if recipient_lines:
        for rl in recipient_lines:
            story.append(rl)
        story.append(Spacer(1, 0.6*cm))

    # ── Date (right-aligned) ───────────────────────────────────────────────────
    ort   = s.get('ort', '').strip()
    today = datetime.now().strftime('%d.%m.%Y')
    date_str = f"{ort + ', ' if ort else ''}{today}"
    story.append(Paragraph(date_str, PRight))
    story.append(Spacer(1, 0.7*cm))

    # ── Subject line ───────────────────────────────────────────────────────────
    if b.get('stelle'):
        story.append(Paragraph(f"Bewerbung als {b['stelle']}", PSubj))
        story.append(Spacer(1, 0.6*cm))

    # ── Body ───────────────────────────────────────────────────────────────────
    anschreiben_text = (b.get('anschreiben') or '').strip()
    if anschreiben_text:
        # Split on 2+ consecutive newlines, keeping the separator to measure blank lines
        parts = re.split(r'(\n{2,})', anschreiben_text)
        for i, part in enumerate(parts):
            if i % 2 == 0:          # text block
                text = part.strip('\n')
                if text:
                    story.append(Paragraph(text.replace('\n', '<br/>'), P))
            else:                    # blank-line separator
                # \n\n = 1 blank line → 0.5cm; each extra \n adds 0.5cm
                extra_lines = len(part) - 2
                story.append(Spacer(1, (0.5 + extra_lines * 0.5) * cm))
    else:
        story.append(Paragraph('<i>Kein Anschreiben hinterlegt.</i>',
                               _style('PI', textColor=GREY, fontSize=11)))

    doc.build(story)
    buf.seek(0)
    return buf


# ── Auth Routes ───────────────────────────────────────────────────────────────

@app.route('/login', methods=['GET', 'POST'])
def login():
    if current_user.is_authenticated:
        return redirect(url_for('index'))
    error = None
    if request.method == 'POST':
        email    = request.form.get('email', '').strip().lower()
        password = request.form.get('password', '')
        conn = get_db()
        row = conn.execute('SELECT * FROM users WHERE email = ?', (email,)).fetchone()
        conn.close()
        if row and bcrypt.check_password_hash(row['password_hash'], password):
            user = User(row['id'], row['email'], row['name'], row['plan'])
            login_user(user, remember=True)
            return redirect(request.args.get('next') or url_for('index'))
        error = 'E-Mail oder Passwort falsch.'
    return render_template('login.html', error=error)


@app.route('/register', methods=['GET', 'POST'])
def register():
    if current_user.is_authenticated:
        return redirect(url_for('index'))
    error = None
    if request.method == 'POST':
        name     = request.form.get('name', '').strip()
        email    = request.form.get('email', '').strip().lower()
        password = request.form.get('password', '')
        if not name or not email or not password:
            error = 'Bitte alle Felder ausfüllen.'
        elif len(password) < 8:
            error = 'Passwort muss mindestens 8 Zeichen haben.'
        else:
            pw_hash = bcrypt.generate_password_hash(password).decode('utf-8')
            conn = get_db()
            try:
                if _PG:
                    cur = conn.execute(
                        'INSERT INTO users (email, password_hash, name) VALUES (?,?,?) RETURNING id',
                        (email, pw_hash, name)
                    )
                    user_id = cur.fetchone()['id']
                else:
                    cur = conn.execute(
                        'INSERT INTO users (email, password_hash, name) VALUES (?,?,?)',
                        (email, pw_hash, name)
                    )
                    user_id = cur.lastrowid
                conn.commit()
                conn.close()
                login_user(User(user_id, email, name), remember=True)
                return redirect(url_for('index'))
            except Exception:
                conn.close()
                error = 'Diese E-Mail-Adresse ist bereits registriert.'
    return render_template('register.html', error=error)


@app.route('/logout')
@login_required
def logout():
    logout_user()
    return redirect(url_for('login'))


if __name__ == '__main__':
    print('\n  BewerbungsKI laeuft auf  http://localhost:5000\n')
    app.run(debug=False, port=5000)
